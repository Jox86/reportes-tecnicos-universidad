# apps/exports/word.py
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x1E, 0x3A, 0x5F)


def _sombrear_celda(celda, color_hex):
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = tc_pr.makeelement(
        qn("w:shd"),
        {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): color_hex},
    )
    tc_pr.append(shd)


def _fila_meta(tabla, etiqueta, valor):
    fila = tabla.add_row().cells
    fila[0].text = etiqueta
    fila[0].paragraphs[0].runs[0].bold = True
    fila[0].paragraphs[0].runs[0].font.color.rgb = AZUL
    fila[1].text = str(valor) if valor not in (None, "") else "—"
    _sombrear_celda(fila[0], "F2F4F7")
    return fila


def generar_word_reporte(reporte) -> BytesIO:
    doc = Document()

    estilo_normal = doc.styles["Normal"]
    estilo_normal.font.name = "Calibri"
    estilo_normal.font.size = Pt(10.5)

    encabezado = doc.add_paragraph("Sistema de Reportes Técnicos — Universidad")
    encabezado.runs[0].font.size = Pt(10)
    encabezado.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    titulo = doc.add_heading(f"Reporte {reporte.codigo}", level=1)
    titulo.runs[0].font.color.rgb = AZUL

    sub = doc.add_paragraph(reporte.tipo_tarea.nombre if reporte.tipo_tarea else "")
    sub.runs[0].font.size = Pt(13)
    sub.runs[0].bold = True

    tabla = doc.add_table(rows=0, cols=2)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.style = "Table Grid"
    tabla.columns[0].width = Cm(4.5)
    tabla.columns[1].width = Cm(11)

    _fila_meta(tabla, "Estado", reporte.get_estado_display())
    _fila_meta(tabla, "Prioridad", reporte.get_prioridad_display())
    _fila_meta(tabla, "Área", f"{reporte.area.nombre} ({reporte.area.tipo_area.nombre})" if reporte.area else "—")
    _fila_meta(tabla, "Ubicación", reporte.ubicacion or "—")
    _fila_meta(tabla, "Usuario afectado", f"{reporte.usuario_nombre}" + (f" — {reporte.usuario_cargo}" if reporte.usuario_cargo else ""))
    _fila_meta(tabla, "Correo del usuario", reporte.usuario_correo)
    _fila_meta(
        tabla,
        "Técnico asignado",
        (reporte.tecnico_asignado.get_full_name() or reporte.tecnico_asignado.username)
        if reporte.tecnico_asignado
        else "Sin asignar",
    )
    _fila_meta(tabla, "Creado por", reporte.creado_por.get_full_name() or reporte.creado_por.username)
    _fila_meta(tabla, "Fecha de creación", reporte.fecha_creacion.strftime("%d/%m/%Y %H:%M"))
    _fila_meta(
        tabla,
        "Fecha de resolución",
        reporte.fecha_resolucion.strftime("%d/%m/%Y %H:%M") if reporte.fecha_resolucion else None,
    )
    _fila_meta(
        tabla,
        "Tiempo de resolución",
        f"{reporte.tiempo_resolucion_horas} h" if reporte.tiempo_resolucion_horas else None,
    )

    # Equipos
    if reporte.equipos.exists():
        h = doc.add_heading("Equipos / activos", level=2)
        h.runs[0].font.color.rgb = AZUL
        tabla2 = doc.add_table(rows=1, cols=4)
        tabla2.style = "Table Grid"
        encabezados = ["Código", "Marca", "Modelo", "N.º de serie"]
        for i, texto in enumerate(encabezados):
            celda = tabla2.rows[0].cells[i]
            celda.text = texto
            celda.paragraphs[0].runs[0].bold = True
            _sombrear_celda(celda, "1E3A5F")
            celda.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for eq in reporte.equipos.all():
            fila = tabla2.add_row().cells
            fila[0].text = eq.codigo_activo or "—"
            fila[1].text = eq.marca or "—"
            fila[2].text = eq.modelo or "—"
            fila[3].text = eq.serie or "—"

    h = doc.add_heading("Descripción", level=2)
    h.runs[0].font.color.rgb = AZUL
    doc.add_paragraph(reporte.descripcion)

    h = doc.add_heading("Solución / notas técnicas", level=2)
    h.runs[0].font.color.rgb = AZUL
    doc.add_paragraph(reporte.solucion or "Sin registrar aún.")

    if reporte.historial.exists():
        h = doc.add_heading("Historial de estados", level=2)
        h.runs[0].font.color.rgb = AZUL
        tabla3 = doc.add_table(rows=1, cols=5)
        tabla3.style = "Table Grid"
        encabezados = ["Fecha", "De", "A", "Usuario", "Comentario"]
        for i, texto in enumerate(encabezados):
            celda = tabla3.rows[0].cells[i]
            celda.text = texto
            celda.paragraphs[0].runs[0].bold = True
            _sombrear_celda(celda, "1E3A5F")
            celda.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for reg in reporte.historial.all().order_by("fecha"):
            fila = tabla3.add_row().cells
            fila[0].text = reg.fecha.strftime("%d/%m/%Y %H:%M")
            fila[1].text = reg.estado_anterior or "—"
            fila[2].text = reg.estado_nuevo
            fila[3].text = (reg.usuario.get_full_name() or reg.usuario.username) if reg.usuario else "—"
            fila[4].text = reg.comentario or "—"

    pie = doc.add_paragraph()
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pie.add_run("Documento generado automáticamente por el Sistema de Reportes Técnicos.")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer