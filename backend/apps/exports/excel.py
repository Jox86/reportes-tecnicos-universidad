# apps/exports/excel.py
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(start_color="1e3a5f", end_color="1e3a5f", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _escribir_encabezados(ws, columnas):
    for col_num, columna in enumerate(columnas, 1):
        cell = ws.cell(row=1, column=col_num, value=columna)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _ajustar_ancho(ws):
    for column_cells in ws.columns:
        length = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(length + 4, 40)


def _excel_safe(value):
    """Evita Excel Formula Injection al exportar texto controlado por usuarios."""
    if not isinstance(value, str):
        return value
    value = value.replace("\x00", "")
    if value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _safe_row(values):
    return [_excel_safe(value) for value in values]


def generar_excel_reportes(queryset):
    """Exporta la lista de reportes a Excel."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reportes"

    columnas = [
        "Código", "Fecha Creación", "Tipo de Tarea", "Área", "Ubicación",
        "Usuario Afectado", "Correo", "Cargo",
        "Prioridad", "Estado", "Técnico Asignado",
        "Fecha Resolución", "Tiempo (h)", "Descripción", "Solución",
    ]
    _escribir_encabezados(ws, columnas)

    for r in queryset:
        ws.append(_safe_row([
            r.codigo,
            r.fecha_creacion.strftime("%Y-%m-%d %H:%M") if r.fecha_creacion else "",
            r.tipo_tarea.nombre if r.tipo_tarea else "",
            r.area.nombre if r.area else "",
            r.ubicacion or "",
            r.usuario_nombre or "",
            r.usuario_correo or "",
            r.usuario_cargo or "",
            r.get_prioridad_display(),
            r.get_estado_display(),
            r.tecnico_nombre or "Sin asignar",
            r.fecha_resolucion.strftime("%Y-%m-%d %H:%M") if r.fecha_resolucion else "",
            r.tiempo_resolucion_horas or "",
            r.descripcion or "",
            r.solucion or "",
        ]))

    _ajustar_ancho(ws)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generar_excel_reporte_individual(reporte):
    """Exporta un reporte individual con formato detallado."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Reporte {reporte.codigo}"

    ws.append(["Campo", "Valor"])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    filas = [
        ("Código", reporte.codigo),
        ("Tipo de Tarea", reporte.tipo_tarea.nombre if reporte.tipo_tarea else ""),
        ("Área", reporte.area.nombre if reporte.area else ""),
        ("Ubicación", reporte.ubicacion or ""),
        ("Usuario Afectado", reporte.usuario_nombre or ""),
        ("Correo", reporte.usuario_correo or ""),
        ("Cargo", reporte.usuario_cargo or ""),
        ("Prioridad", reporte.get_prioridad_display()),
        ("Estado", reporte.get_estado_display()),
        ("Técnico Asignado", reporte.tecnico_nombre or "Sin asignar"),
        ("Creado por", reporte.creado_por_nombre),
        ("Fecha Creación", reporte.fecha_creacion.strftime("%Y-%m-%d %H:%M") if reporte.fecha_creacion else ""),
        ("Fecha Resolución", reporte.fecha_resolucion.strftime("%Y-%m-%d %H:%M") if reporte.fecha_resolucion else ""),
        ("Tiempo (h)", reporte.tiempo_resolucion_horas or ""),
        ("Descripción", reporte.descripcion or ""),
        ("Solución", reporte.solucion or ""),
    ]
    for fila in filas:
        ws.append(_safe_row(fila))

    if reporte.equipos.exists():
        ws.append([])
        ws.append(["Equipos", ""])
        for cell in ws[ws.max_row]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        ws.append(["Código", "Marca", "Modelo", "Serie"])
        for eq in reporte.equipos.all():
            ws.append(_safe_row([eq.codigo_activo, eq.marca, eq.modelo, eq.serie]))

    _ajustar_ancho(ws)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
